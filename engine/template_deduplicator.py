"""Conservative static-skeleton fingerprinting for mined DOCX templates."""
from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Iterable

from docx import Document

PLACEHOLDER_RE = re.compile(r"\{[^{}]+\}")
SPACE_RE = re.compile(r"\s+")
PUNCTUATION_RE = re.compile(r"[\W_]+", re.UNICODE)

@dataclass(frozen=True)
class Fingerprint:
    path: str
    digest: str
    skeleton: str
    character_count: int

@dataclass(frozen=True)
class DuplicateMatch:
    candidate_path: str
    matched_path: str
    similarity: float
    exact_hash: bool
    decision: str


def normalize_skeleton(text: str) -> str:
    """Normalize Unicode and whitespace/punctuation without deleting words."""
    text = unicodedata.normalize("NFKC", text)
    text = PLACEHOLDER_RE.sub(" <FIELD> ", text)
    text = PUNCTUATION_RE.sub(" ", text)
    return SPACE_RE.sub(" ", text).strip().casefold()


def _document_text(path: str | Path) -> str:
    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def fingerprint_docx(path: str | Path) -> Fingerprint:
    path = Path(path)
    skeleton = normalize_skeleton(_document_text(path))
    digest = hashlib.sha256(skeleton.encode("utf-8")).hexdigest()
    return Fingerprint(str(path), digest, skeleton, len(skeleton))


def compare_fingerprints(left: Fingerprint, right: Fingerprint) -> float:
    """Return a conservative similarity score in [0, 1].

    SequenceMatcher is useful for review ranking, but should not be interpreted
    as proof that legal text is interchangeable.
    """
    if not left.skeleton and not right.skeleton:
        return 1.0
    if not left.skeleton or not right.skeleton:
        return 0.0
    return SequenceMatcher(None, left.skeleton, right.skeleton, autojunk=False).ratio()


def classify_duplicate(candidate: Fingerprint, existing: Iterable[Fingerprint], threshold: float = 0.95) -> DuplicateMatch | None:
    """Return the strongest likely duplicate; never silently delete a file."""
    if not 0.0 <= threshold <= 1.0:
        raise ValueError("threshold must be between 0.0 and 1.0")
    best: DuplicateMatch | None = None
    for known in existing:
        exact = candidate.digest == known.digest
        score = 1.0 if exact else compare_fingerprints(candidate, known)
        if score < threshold:
            continue
        decision = "exact_duplicate" if exact else "review_near_duplicate"
        match = DuplicateMatch(candidate.path, known.path, score, exact, decision)
        if best is None or match.similarity > best.similarity:
            best = match
    return best


def scan_docx_folder(folder_path: str | Path, threshold: float = 0.95) -> tuple[list[Fingerprint], list[DuplicateMatch]]:
    """Fingerprint DOCX files and group likely duplicates, preserving all files."""
    folder = Path(folder_path)
    if not folder.is_dir():
        raise NotADirectoryError(folder)
    unique: list[Fingerprint] = []
    duplicates: list[DuplicateMatch] = []
    for path in sorted(folder.rglob("*.docx")):
        if path.name.startswith("~$"):
            continue
        candidate = fingerprint_docx(path)
        match = classify_duplicate(candidate, unique, threshold)
        if match is None:
            unique.append(candidate)
        else:
            duplicates.append(match)
    return unique, duplicates
