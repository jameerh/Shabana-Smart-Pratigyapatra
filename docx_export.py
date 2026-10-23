# docx_export.py
import os
import tempfile
import sys
sys.path.append("engine")  # ensure engine folder is in path
try:
    from ism_converter import unicode_to_ism
    ISM_AVAILABLE = True
except ImportError:
    ISM_AVAILABLE = False
    def unicode_to_ism(text):
        return text

from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import barcode
from barcode.writer import ImageWriter

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ['top', 'left', 'bottom', 'right']:
        edge_attr = kwargs.get(edge)
        if edge_attr is not None:
            tag = 'w:{}'.format(edge)
            border = OxmlElement(tag)
            border.set(qn('w:val'), edge_attr.get('val', 'single'))
            border.set(qn('w:sz'), edge_attr.get('sz', '4'))
            border.set(qn('w:space'), edge_attr.get('space', '0'))
            border.set(qn('w:color'), edge_attr.get('color', '000000'))
            tcPr.append(border)

def add_page_border(doc):
    section = doc.sections[0]
    sectPr = section._sectPr
    borders = OxmlElement('w:pgBorders')
    for edge in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '24')
        border.set(qn('w:color'), '000000')
        borders.append(border)
    sectPr.append(borders)

def add_stamp_box(doc, font_name, use_ism):
    """Add a 2x2 inch box for stamp, with appropriate font and conversion."""
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.allow_autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(2)
    cell.height = Inches(2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell, top={'val':'single','sz':'4'}, left={'val':'single','sz':'4'},
                    bottom={'val':'single','sz':'4'}, right={'val':'single','sz':'4'})
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stamp_text = "स्टॅम्पसाठी जागा"
    if use_ism and ISM_AVAILABLE:
        stamp_text = unicode_to_ism(stamp_text)
    run = p.add_run(stamp_text)
    run.font.name = font_name
    run.font.size = Pt(10)
    run.italic = True
    return table

def add_header(doc, taluka_name, font_name, use_ism):
    p = doc.add_paragraph()
    header_text = f"समक्ष: तहसीलदार, {taluka_name if taluka_name else '________'}"
    if use_ism and ISM_AVAILABLE:
        header_text = unicode_to_ism(header_text)
    run = p.add_run(header_text)
    run.font.name = font_name
    run.font.size = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)

def add_notary_block(doc, font_name, use_ism):
    doc.add_paragraph()
    # शपथपत्रकर्ता
    p = doc.add_paragraph()
    text = "शपथपत्रकर्ता"
    if use_ism and ISM_AVAILABLE:
        text = unicode_to_ism(text)
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(12)
    # स्वाक्षरी
    p = doc.add_paragraph()
    text = "स्वाक्षरी"
    if use_ism and ISM_AVAILABLE:
        text = unicode_to_ism(text)
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(6)
    # signature line
    p = doc.add_paragraph()
    run = p.add_run("______________________")
    run.font.name = font_name
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(12)
    # सामोर शपथ दिली
    p = doc.add_paragraph()
    text = "सामोर शपथ दिली"
    if use_ism and ISM_AVAILABLE:
        text = unicode_to_ism(text)
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(6)
    # नोटरी
    p = doc.add_paragraph()
    text = "नोटरी"
    if use_ism and ISM_AVAILABLE:
        text = unicode_to_ism(text)
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(12)
    run.bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(12)

def add_page_number(doc, font_name, use_ism):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_text = "पृष्ठ क्र. ________"
    if use_ism and ISM_AVAILABLE:
        page_text = unicode_to_ism(page_text)
    run = p.add_run(page_text)
    run.font.name = font_name
    run.font.size = Pt(10)

def generate_barcode_image(data, output_path):
    try:
        code128 = barcode.get_barcode_class('code128')
        barcode_obj = code128(data, writer=ImageWriter())
        barcode_obj.save(output_path)
        return output_path + ".png"
    except Exception as e:
        print(f"Barcode generation failed: {e}")
        return None

def add_barcode(doc, barcode_data, position="top-right"):
    if not barcode_data:
        return
    temp_file = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    temp_file.close()
    barcode_path = generate_barcode_image(barcode_data, temp_file.name)
    if not barcode_path or not os.path.exists(barcode_path):
        return
    if position == "top-right":
        p = doc.add_paragraph()
        doc.paragraphs.insert(0, p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    run.add_picture(barcode_path, width=Cm(6.5))
    try:
        os.unlink(barcode_path)
    except:
        pass

def combine_paragraphs(lines):
    paragraphs = []
    current = []
    for line in lines:
        stripped = line.strip()
        if stripped == "":
            if current:
                paragraphs.append(" ".join(current))
                current = []
        else:
            current.append(stripped)
    if current:
        paragraphs.append(" ".join(current))
    return paragraphs

def export_to_docx(text, filename="affidavit.docx", barcode_data=None, taluka_name="", use_ism=False):
    # Determine font and optionally convert the main text
    if use_ism and ISM_AVAILABLE:
        text = unicode_to_ism(text)
        font_name = "DVBW"   # or any ISM font installed
    else:
        font_name = "Mangal"

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    add_page_border(doc)
    add_header(doc, taluka_name, font_name, use_ism)
    add_stamp_box(doc, font_name, use_ism)

    if barcode_data:
        add_barcode(doc, barcode_data, position="top-right")

    lines = text.split("\n")
    while lines and lines[0].strip() == "":
        lines.pop(0)
    title_line = lines.pop(0).strip() if lines else "प्रतिज्ञापत्र"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title_line)
    run.font.name = font_name
    run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(14)
    run.bold = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)

    paragraphs = combine_paragraphs(lines)
    for para_text in paragraphs:
        p = doc.add_paragraph()
        run = p.add_run(para_text)
        run.font.name = font_name
        run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        run.font.size = Pt(14)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Inches(0.5)

    add_notary_block(doc, font_name, use_ism)
    add_page_number(doc, font_name, use_ism)

    doc.save(filename)

def export_waras_template(filename="waras_template.docx", taluka_name=""):
    # For blank templates, we keep the text in Unicode (no ISM conversion)
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    add_page_border(doc)
    # For the blank template, we use Mangal font, no conversion
    font_name = "Mangal"
    use_ism = False
    add_header(doc, taluka_name, font_name, use_ism)
    add_stamp_box(doc, font_name, use_ism)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("प्रतिज्ञापत्र")
    run.font.name = font_name
    run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(14)
    run.bold = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)

    body_paragraphs = [
        "मी __________________________, वय ___________ वर्षे, व्यवसाय __________________________,",
        "रा. __________________________, ता. __________________________, जि. __________________________,",
        "",
        "मी सत्यप्रतिज्ञेवर सांगतो/सांगते की,",
        "",
        "माझे __________________________ __________________________ यांचा मृत्यू दिनांक __________",
        "रोजी __________________________ येथे झाला.",
        "",
        "त्यांचे खालीलप्रमाणे वारस आहेत:"
    ]
    for para_text in body_paragraphs:
        p = doc.add_paragraph()
        run = p.add_run(para_text)
        run.font.name = font_name
        run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        run.font.size = Pt(14)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Inches(0.5)

    rows = 6
    table = doc.add_table(rows=rows, cols=3)
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        row.cells[0].width = Inches(1.2)
        row.cells[1].width = Inches(3.0)
        row.cells[2].width = Inches(2.5)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "अ.क्र."
    hdr_cells[1].text = "वारसाचे नाव"
    hdr_cells[2].text = "वय (वर्षे)"
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            run.font.name = font_name
            run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
            run.font.size = Pt(12)
            run.bold = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for i in range(1, rows):
        row_cells = table.rows[i].cells
        row_cells[0].text = str(i)
        row_cells[1].text = "__________________________"
        row_cells[2].text = "________"
        for cell in row_cells:
            for paragraph in cell.paragraphs:
                run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                run.font.name = font_name
                run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
                run.font.size = Pt(12)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    rest_paragraphs = [
        "वरीलप्रमाणे आम्ही _____ कायदेशीर वारस असून आमच्याशिवाय अन्य कोणीही वारस नाही.",
        "",
        "तरी __________________________ या मालमत्तेस वारसनोंद करण्यासाठी सदर प्रतिज्ञापत्र देत आहे.",
        "",
        "ठिकाण: __________________________",
        "दिनांक: __________________________",
        "",
        "स्वाक्षरी",
        "__________________________"
    ]
    for para_text in rest_paragraphs:
        p = doc.add_paragraph()
        run = p.add_run(para_text)
        run.font.name = font_name
        run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        run.font.size = Pt(14)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Inches(0.5)

    add_notary_block(doc, font_name, use_ism)
    add_page_number(doc, font_name, use_ism)

    doc.save(filename)

def export_waras_filled(filename, data, taluka_name="", use_ism=False):
    """
    Create a filled Waras affidavit with a table for heirs.
    """
    if use_ism and ISM_AVAILABLE:
        font_name = "DVBW"
        # Convert the data values that are text
        for key, val in data.items():
            if isinstance(val, str):
                data[key] = unicode_to_ism(val)
    else:
        font_name = "Mangal"

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    add_page_border(doc)
    add_header(doc, taluka_name, font_name, use_ism)
    add_stamp_box(doc, font_name, use_ism)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("प्रतिज्ञापत्र")
    run.font.name = font_name
    run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(14)
    run.bold = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)

    # Body paragraphs
    body_paragraphs = [
        f"मी {data['applicant_full_name']}, वय {data['applicant_age']} वर्षे, व्यवसाय {data['applicant_occupation']},",
        f"रा. {data['applicant_taluka']}, ता. {data['applicant_taluka']}, जि. {data['applicant_district']},",
        "",
        f"मी सत्यप्रतिज्ञेवर {data.get('gender_verb', 'सांगतो')} की,",
        "",
        f"माझे {data['deceased_relation']} {data['deceased_name']} यांचा मृत्यू दिनांक {data['deceased_death_date']}",
        f"रोजी {data['deceased_death_place']} येथे झाला.",
        "",
        "त्यांचे खालीलप्रमाणे वारस आहेत:"
    ]
    for para_text in body_paragraphs:
        if para_text.strip() == "":
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        run = p.add_run(para_text)
        run.font.name = font_name
        run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        run.font.size = Pt(14)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Inches(0.5)

    # Heir table
    heirs = data.get("heirs_list", [])
    if heirs:
        rows = len(heirs) + 1  # +1 for header
        table = doc.add_table(rows=rows, cols=4)
        table.autofit = False
        table.allow_autofit = False
        # Set column widths (inches)
        col_widths = [0.5, 3.0, 1.0, 1.5]
        for i, width in enumerate(col_widths):
            table.columns[i].width = Inches(width)

        # Header row
        hdr_cells = table.rows[0].cells
        headers = ["अ.क्र.", "वारसाचे नाव", "वय", "नाते"]
        for i, text in enumerate(headers):
            cell = hdr_cells[i]
            cell.text = text
            for paragraph in cell.paragraphs:
                run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                run.font.name = font_name
                run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
                run.font.size = Pt(12)
                run.bold = True
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Data rows
        for i, (name, age, rel) in enumerate(heirs, start=1):
            row_cells = table.rows[i].cells
            row_cells[0].text = str(i)
            row_cells[1].text = name
            row_cells[2].text = age
            row_cells[3].text = rel
            for cell in row_cells:
                for paragraph in cell.paragraphs:
                    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                    run.font.name = font_name
                    run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
                    run.font.size = Pt(12)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Remaining text
    rest_paragraphs = [
        "",
        f"वरीलप्रमाणे आम्ही {data['number_of_heirs']} कायदेशीर वारस असून आमच्याशिवाय अन्य कोणीही वारस नाही.",
        "",
        f"तरी {data['city_survey_number']} या मालमत्तेस वारसनोंद करण्यासाठी सदर प्रतिज्ञापत्र देत आहे.",
        "",
        f"ठिकाण: {data['affidavit_place']}",
        f"दिनांक: {data['affidavit_date']}",
        "",
        "स्वाक्षरी",
        f"({data['applicant_full_name']})"
    ]
    for para_text in rest_paragraphs:
        if para_text.strip() == "":
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        run = p.add_run(para_text)
        run.font.name = font_name
        run.font._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        run.font.size = Pt(14)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Inches(0.5)

    add_notary_block(doc, font_name, use_ism)
    add_page_number(doc, font_name, use_ism)

    doc.save(filename)