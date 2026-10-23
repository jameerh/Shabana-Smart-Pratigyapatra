"""Unit test for intra-document repeating section detection and batch pipeline."""
import tempfile
import unittest
from pathlib import Path
from docx import Document

from engine.section_repeater import detect_repeated_person_sections
from engine.batch_pipeline import process_template_folder


class BatchPipelineTests(unittest.TestCase):
    def test_detects_repeated_person_sections(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc_path = Path(tmp) / "multi_person.docx"
            doc = Document()
            # Page 1
            doc.add_paragraph("प्रतिज्ञापत्र")
            p1 = doc.add_paragraph("मी ")
            p1.add_run("रमेश पाटील").bold = True
            p1.add_run(", वय ")
            p1.add_run("४५").bold = True
            p1.add_run(" वर्षे, रा. पुणे.")
            doc.add_paragraph("सत्यप्रतिज्ञेवर सांगतो की सर्व माहिती खरी आहे.")

            # Page 2 (Repeated layout for Dharmesh)
            doc.add_paragraph("प्रतिज्ञापत्र")
            p2 = doc.add_paragraph("मी ")
            p2.add_run("धर्मेश पाटील").bold = True
            p2.add_run(", वय ")
            p2.add_run("३८").bold = True
            p2.add_run(" वर्षे, रा. पुणे.")
            doc.add_paragraph("सत्यप्रतिज्ञेवर सांगतो की सर्व माहिती खरी आहे.")
            doc.save(doc_path)

            result = detect_repeated_person_sections(doc_path)
            self.assertTrue(result.is_repeated)
            self.assertGreaterEqual(result.segment_count, 2)

    def test_process_template_folder_deduplication(self):
        with tempfile.TemporaryDirectory() as in_dir, tempfile.TemporaryDirectory() as out_dir:
            # File 1: Ramesh
            f1 = Path(in_dir) / "affidavit_ramesh.docx"
            doc1 = Document()
            p1 = doc1.add_paragraph("अर्जदाराचे नाव: ")
            p1.add_run("रमेश पाटील").bold = True
            doc1.add_paragraph("मी सत्यप्रतिज्ञेवर सांगतो की मी पुणे येथे राहतो.")
            doc1.save(f1)

            # File 2: Same template filled for Suresh (Duplicate skeleton)
            f2 = Path(in_dir) / "affidavit_suresh.docx"
            doc2 = Document()
            p2 = doc2.add_paragraph("अर्जदाराचे नाव: ")
            p2.add_run("सुरेश शिंदे").bold = True
            doc2.add_paragraph("मी सत्यप्रतिज्ञेवर सांगतो की मी पुणे येथे राहतो.")
            doc2.save(f2)

            # File 3: Different template (Income)
            f3 = Path(in_dir) / "income_declaration.docx"
            doc3 = Document()
            p3 = doc3.add_paragraph("माझे वार्षिक उत्पन्न: ")
            p3.add_run("५०,०००").bold = True
            doc3.add_paragraph("तसेच मी शेतमजूर म्हणून काम करतो.")
            doc3.save(f3)

            report = process_template_folder(in_dir, out_dir, similarity_threshold=0.90)
            self.assertEqual(report.total_scanned, 3)
            self.assertEqual(report.unique_mined, 2)
            self.assertEqual(report.duplicates_skipped, 1)
            self.assertEqual(len(report.skipped_files), 1)
            self.assertEqual(report.skipped_files[0]["file"], "affidavit_suresh.docx")


if __name__ == "__main__":
    unittest.main()
