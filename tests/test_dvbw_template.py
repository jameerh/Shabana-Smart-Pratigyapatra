"""Regression tests for the DVBW legacy-font template and font detection."""
import tempfile
import unittest
from pathlib import Path

from docx import Document

from templates import TEMPLATES
from engine.smart_docx_import import detect_docx_font_family, extract_docx_fields, build_docx_preview


class DvbwTemplateTests(unittest.TestCase):
    def test_dvbw_affidavit_is_registered_and_uses_placeholders(self):
        info = TEMPLATES["DVBW DDF Electricity Connection Affidavit"]
        self.assertIn("applicant_full_name", info["fields"])
        self.assertIn("{applicant_full_name}", info["template"])
        self.assertIn("{property_number}", info["template"])
        self.assertEqual(info["preview_font_hint"], "DVBW")

    def test_preview_replaces_dvbw_template_placeholders(self):
        info = TEMPLATES["DVBW DDF Electricity Connection Affidavit"]
        rendered = info["template"].format(
            applicant_full_name="टेस्ट नाव", applicant_age="40",
            applicant_occupation="शेती", applicant_village="जत",
            property_number="473/1/A", connection_percentage="1.3%",
            electricity_office="जत", affidavit_place="जत",
            affidavit_date="23/10/2026",
        )
        self.assertIn("टेस्ट नाव", rendered)
        self.assertIn("473/1/A", rendered)
        self.assertNotIn("{applicant_full_name}", rendered)

    def test_detects_explicit_font_and_placeholder_in_docx(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "legacy.docx"
            doc = Document()
            run = doc.add_paragraph().add_run("नाव {applicant_full_name}")
            run.font.name = "DVBW Series"
            doc.save(path)
            self.assertEqual(detect_docx_font_family(path), "DVBW Series")
            self.assertEqual(extract_docx_fields(path)[0]["key"], "applicant_full_name")
            preview = build_docx_preview(path, {"applicant_full_name": "राम"})
            self.assertIn("राम", preview)


if __name__ == "__main__":
    unittest.main()
