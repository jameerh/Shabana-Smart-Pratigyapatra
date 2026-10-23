import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt

from engine.bold_template_miner import mine_docx


class BoldTemplateMinerTests(unittest.TestCase):
    def test_merges_adjacent_bold_runs_and_replaces_them(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "legacy.docx"
            output = Path(tmp) / "mined.docx"
            doc = Document()
            p = doc.add_paragraph("अर्जदाराचे नाव: ")
            p.add_run("रमेश").bold = True
            p.add_run(" ").bold = True
            p.add_run("पाटील").bold = True
            doc.save(source)

            result = mine_docx(source, output)
            self.assertEqual(len(result.fields), 1)
            self.assertEqual(result.fields[0].value, "रमेश पाटील")
            self.assertEqual(result.fields[0].name, "applicant_name")
            mined = Document(output)
            self.assertIn("{applicant_name}", mined.paragraphs[0].text)
            self.assertEqual(Document(source).paragraphs[0].text, "अर्जदाराचे नाव: रमेश पाटील")

    def test_does_not_mine_regular_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "plain.docx"
            doc = Document()
            doc.add_paragraph("नाव: रमेश पाटील")
            doc.save(source)
            result = mine_docx(source)
            self.assertEqual(result.fields, [])
            self.assertTrue(result.warnings)

    def test_handles_bold_text_in_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "table.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "नाव"
            run = table.cell(0, 1).paragraphs[0].add_run("सावित्री")
            run.bold = True
            doc.save(source)
            result = mine_docx(source)
            self.assertEqual(len(result.fields), 1)
            self.assertEqual(result.fields[0].value, "सावित्री")


if __name__ == "__main__":
    unittest.main()
