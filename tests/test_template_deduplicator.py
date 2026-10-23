import tempfile
import unittest
from pathlib import Path

from docx import Document

from engine.template_deduplicator import (
    classify_duplicate,
    fingerprint_docx,
    normalize_skeleton,
)


class TemplateDeduplicatorTests(unittest.TestCase):
    def test_normalization_ignores_punctuation_and_case(self):
        self.assertEqual(normalize_skeleton("Name: {applicant_name}, Age."),
                         normalize_skeleton("name {person} age"))

    def test_exact_skeleton_detected_after_values_become_placeholders(self):
        with tempfile.TemporaryDirectory() as tmp:
            first = Path(tmp) / "first.docx"
            second = Path(tmp) / "second.docx"
            for path, name in ((first, "{applicant_name}"), (second, "{person_name}")):
                doc = Document()
                doc.add_paragraph(f"मी {name} सत्यप्रतिज्ञेवर सांगतो की")
                doc.save(path)
            a = fingerprint_docx(first)
            b = fingerprint_docx(second)
            self.assertEqual(a.digest, b.digest)
            match = classify_duplicate(b, [a])
            self.assertIsNotNone(match)
            self.assertTrue(match.exact_hash)
            self.assertEqual(match.decision, "exact_duplicate")

    def test_empty_documents_are_not_crash_cases(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "empty.docx"
            Document().save(path)
            fp = fingerprint_docx(path)
            self.assertEqual(fp.character_count, 0)


if __name__ == "__main__":
    unittest.main()
