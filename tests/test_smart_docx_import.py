"""Tests for automatic placeholder detection and DOCX filling."""
import tempfile
import unittest
from pathlib import Path

from docx import Document

from engine.smart_docx_import import extract_docx_fields, build_docx_preview
from template_parser import create_filled_document


class SmartDocxImportTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.source = self.root / "sample.docx"
        self.output = self.root / "filled.docx"

        doc = Document()
        doc.add_paragraph("प्रतिज्ञापत्र: [प्रतिज्ञाकर्त्याचे पूर्ण नाव], वय [वय]")
        doc.add_paragraph("गाव: {village}")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "तालुका [तालुका]"
        doc.save(self.source)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_detects_square_and_brace_placeholders_in_document_order(self):
        fields = extract_docx_fields(self.source)
        self.assertEqual(
            [field["key"] for field in fields],
            ["प्रतिज्ञाकर्त्याचे पूर्ण नाव", "वय", "village", "तालुका"],
        )
        self.assertEqual(fields[1]["label"], "वय (वर्षे)")

    def test_preview_replaces_filled_values_and_keeps_empty_placeholders(self):
        preview = build_docx_preview(
            self.source,
            {"प्रतिज्ञाकर्त्याचे पूर्ण नाव": "राम पाटील", "village": "जत"},
        )
        self.assertIn("राम पाटील", preview)
        self.assertIn("जत", preview)
        self.assertIn("[वय]", preview)
        self.assertIn("[तालुका]", preview)

    def test_export_replaces_placeholders_in_paragraphs_and_tables(self):
        create_filled_document(
            self.source,
            {
                "प्रतिज्ञाकर्त्याचे पूर्ण नाव": "राम पाटील",
                "वय": "35",
                "village": "जत",
                "तालुका": "जत",
            },
            self.output,
        )
        result = Document(self.output)
        self.assertIn("राम पाटील", "\\n".join(p.text for p in result.paragraphs))
        self.assertIn("35", "\\n".join(p.text for p in result.paragraphs))
        self.assertIn("जत", result.tables[0].cell(0, 0).text)

    def test_preview_docx_keeps_empty_placeholders_but_export_removes_them(self):
        create_filled_document(
            self.source,
            {"प्रतिज्ञाकर्त्याचे पूर्ण नाव": "राम पाटील"},
            self.output,
            show_empty=True,
        )
        preview_doc = Document(self.output)
        preview_text = "\\n".join(p.text for p in preview_doc.paragraphs)
        self.assertIn("[वय]", preview_text)
        self.assertIn("{village}", preview_text)

        create_filled_document(
            self.source,
            {"प्रतिज्ञाकर्त्याचे पूर्ण नाव": "राम पाटील"},
            self.output,
            show_empty=False,
        )
        export_doc = Document(self.output)
        export_text = "\\n".join(p.text for p in export_doc.paragraphs)
        self.assertNotIn("[वय]", export_text)
        self.assertNotIn("{village}", export_text)

    def test_replaces_placeholder_split_across_word_runs(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("नाव: [applicant_")
        paragraph.add_run("full_name]")
        doc.save(self.source)

        create_filled_document(
            self.source,
            {"applicant_full_name": "सीमा पाटील"},
            self.output,
        )
        result = Document(self.output)
        self.assertEqual(result.paragraphs[0].text, "नाव: सीमा पाटील")


if __name__ == "__main__":
    unittest.main()
